"""Build the reusable compact, blue-heading daily briefing document."""
import argparse
import json
import re
import sys
import zipfile
from pathlib import Path
from PIL import Image
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
sys.dont_write_bytecode=True
from font_embedding import embed_subset


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input',required=True)
    parser.add_argument('--output',required=True)
    parser.add_argument('--font',required=True)
    parser.add_argument('--font-family',default='Noto Sans CJK SC')
    parser.add_argument('--fonttools-wheel')
    args=parser.parse_args()
    source=Path(args.input).resolve();output=Path(args.output).resolve()
    data=json.loads(source.read_text());doc=Document();sec=doc.sections[0]
    if not data.get('title') or not re.fullmatch(r'\d{4}-\d{2}-\d{2}',data.get('publication_date','')):
        raise ValueError('Article title and publication_date are required')
    sec.orientation=WD_ORIENT.LANDSCAPE;sec.page_width=Cm(29.7);sec.page_height=Cm(21)
    sec.left_margin=Cm(1.5);sec.right_margin=Cm(1.5);sec.top_margin=Cm(1.25);sec.bottom_margin=Cm(1.25)
    sec.header_distance=Cm(.5);sec.footer_distance=Cm(.5)
    blue='2A77BF';family=args.font_family

    def style(name,size,bold=False,color='000000'):
        target=doc.styles[name];target.font.name=family;target.font.size=Pt(size)
        target.font.bold=bold;target.font.color.rgb=RGBColor.from_string(color)
        fonts=target.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ['ascii','hAnsi','eastAsia']:fonts.set(qn('w:'+attr),family)
        color_node=target.element.get_or_add_rPr().find(qn('w:color'))
        for attr in ['themeColor','themeTint','themeShade']:color_node.attrib.pop(qn('w:'+attr),None)
        para=target.paragraph_format;para.space_before=Pt(0);para.space_after=Pt(6);para.line_spacing=1.15
        ppr=target.element.get_or_add_pPr();snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');ppr.append(snap)
        for border in list(target.element.xpath('.//w:pBdr')):border.getparent().remove(border)
        return para

    style('Normal',11)
    title=style('Title',21,True,blue);title.keep_with_next=True;title.space_after=Pt(5)
    sub=style('Subtitle',10);sub.keep_with_next=True;sub.space_after=Pt(9)
    for name,size in [('Heading 1',15),('Heading 2',12)]:
        heading=style(name,size,False,blue);heading.keep_with_next=True;heading.space_before=Pt(8);heading.space_after=Pt(5)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('第 ').font.size=Pt(8);field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field);footer.add_run(' 页').font.size=Pt(8)
    date=data['publication_date'];year,month,day=map(int,date.split('-'))
    header=sec.header.paragraphs[0];header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    header.add_run(f'NBA Fantasy 每日简报  {year}年{month}月{day}日').font.size=Pt(8)
    emoji_files=data.get('emoji_files',{})
    splitter='('+'|'.join(re.escape(k) for k in sorted(emoji_files,key=len,reverse=True))+')' if emoji_files else None

    def add_text(p,text,bold=None):
        for piece in re.split(splitter,str(text)) if splitter else [str(text)]:
            if not piece:continue
            run=p.add_run()
            if piece in emoji_files:
                picture=run.add_picture(str(Path(emoji_files[piece])),width=Cm(.34));picture._inline.docPr.set('descr',piece)
            else:
                run.text=piece
                if bold is not None:run.bold=bold

    def paragraph(text,label=None):
        p=doc.add_paragraph();p.paragraph_format.keep_together=True
        if label:add_text(p,label+'　',True)
        add_text(p,text);return p

    add_text(doc.add_paragraph(style='Title'),data['title'])
    add_text(doc.add_paragraph(style='Subtitle'),data.get('subtitle') or f'{year}年{month}月{day}日')
    expected_cards=[]
    for section in data.get('sections',[]):
        if section.get('title'):add_text(doc.add_paragraph(style='Heading 1'),section['title'])
        for text in section.get('paragraphs',[]):paragraph(text)
        for player in section.get('players',[]):
            card=Path(player['card_file']);label=player.get('label') or player.get('chineseName') or player['name']
            with Image.open(card) as image:
                if image.size!=(2400,200):raise ValueError(f'Expected a compact card for {label}')
            p=doc.add_paragraph();p.paragraph_format.keep_with_next=True;p.paragraph_format.keep_together=True
            p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.0
            run=p.add_run();run.font.size=Pt(1);picture=run.add_picture(str(card),width=Cm(26.7))
            s=player['stats']
            picture._inline.docPr.set('title',label+' 当日球员数据卡')
            picture._inline.docPr.set('descr',f"{label} {player['team']} {player['date']} {player['minutes']}分钟 {s['pts']}分{s['reb']}板{s['ast']}助 三分{s['three']} 投篮{s['fgm']}/{s['fga']} 罚球{s['ftm']}/{s['fta']} {s['st']}断{s['blk']}帽{s['to']}失误")
            p=paragraph(player['comment'],label);p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(8)
            expected_cards.append(label)
    if data.get('injuries'):
        add_text(doc.add_paragraph(style='Heading 1'),data.get('injury_heading','一些伤病更新'))
        for item in data['injuries']:paragraph(item['text'],item.get('team'))
    doc.core_properties.title=data['title'];doc.core_properties.subject='NBA Fantasy 带图简报'
    doc.core_properties.author='';doc.core_properties.last_modified_by='';doc.core_properties.comments=''
    output.parent.mkdir(parents=True,exist_ok=True);doc.save(output)
    report=embed_subset(output,Path(args.font),family,args.fonttools_wheel)
    with zipfile.ZipFile(output) as z:
        xml=z.read('word/document.xml').decode();cards=xml.count('当日球员数据卡')
        if cards!=len(expected_cards):raise AssertionError('Player card count mismatch')
        if 'Image</w:t>' in xml:raise AssertionError('Image placeholder remains in the document')
        media=[n for n in z.namelist() if n.startswith('word/media/')]
    size=output.stat().st_size
    if size>=15_000_000:raise ValueError(f'DOCX is {size} bytes and exceeds the 15MB limit')
    report.update(file=str(output),bytes=size,limit_bytes=15_000_000,player_cards=cards,embedded_images=len(media),heading_color='#'+blue,card_dimensions=[2400,200],players=expected_cards)
    output.with_suffix('.verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
    print(json.dumps(report,ensure_ascii=False))


if __name__=='__main__':main()
